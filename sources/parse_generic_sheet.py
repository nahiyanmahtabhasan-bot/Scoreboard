"""Parse arbitrary spreadsheet tabs into columns + rows."""

from __future__ import annotations

import math
from typing import Any, BinaryIO

from openpyxl import load_workbook


def _cell_str(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    text = str(value).strip()
    return text or None


def _cell_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, str):
        return value.strip()
    return value


def list_sheet_names(source: str | BinaryIO) -> list[str]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def parse_generic_sheet(
    source: str | BinaryIO,
    *,
    tab: str | None = None,
    max_rows: int = 500,
    source_label: str | None = None,
) -> dict[str, Any]:
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        names = list(workbook.sheetnames)
        if not names:
            raise ValueError("Spreadsheet has no sheets")
        requested = (tab or "").strip() or None
        if requested:
            if requested not in names:
                raise ValueError(
                    f"Sheet tab {requested!r} not found. Available: {', '.join(names)}"
                )
            sheet_name = requested
        else:
            sheet_name = names[0]
        worksheet = workbook[sheet_name]
        raw_rows = list(worksheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not raw_rows:
        return {
            "kind": "table",
            "source_file": source_label or "Spreadsheet",
            "source_sheet": sheet_name,
            "sheet_names": names,
            "columns": [],
            "rows": [],
        }

    header_idx = 0
    for idx, row in enumerate(raw_rows[:20]):
        non_empty = [_cell_str(c) for c in row if _cell_str(c)]
        if len(non_empty) >= 1:
            header_idx = idx
            break

    header_row = raw_rows[header_idx]
    columns: list[str] = []
    seen: dict[str, int] = {}
    for col_idx, cell in enumerate(header_row):
        name = _cell_str(cell) or f"Column {col_idx + 1}"
        count = seen.get(name, 0)
        seen[name] = count + 1
        columns.append(name if count == 0 else f"{name} ({count + 1})")

    # Trim trailing empty header columns
    while columns and columns[-1].startswith("Column ") and all(
        _cell_value(raw_rows[r][len(columns) - 1]) is None
        for r in range(header_idx + 1, min(len(raw_rows), header_idx + 6))
        if len(raw_rows[r]) >= len(columns)
    ):
        columns.pop()

    rows: list[dict[str, Any]] = []
    for row in raw_rows[header_idx + 1 : header_idx + 1 + max_rows]:
        if all(_cell_value(c) is None for c in row):
            continue
        entry: dict[str, Any] = {}
        empty = True
        for i, col in enumerate(columns):
            val = _cell_value(row[i]) if i < len(row) else None
            entry[col] = val
            if val is not None and val != "":
                empty = False
        if not empty:
            rows.append(entry)

    return {
        "kind": "table",
        "source_file": source_label or "Spreadsheet",
        "source_sheet": sheet_name,
        "sheet_names": names,
        "columns": columns,
        "rows": rows,
        "has_scorecard_tab": "Scorecard" in names,
    }
